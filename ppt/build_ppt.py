"""Builds SkyTrace_Presentation.pptx (with speaker notes) from content.py.
Import into Canva: Create a design -> Import file."""
import json, os, sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "ppt")
sys.path.insert(0, OUT)
from content import SLIDES

D = json.load(open(os.path.join(ROOT, "web", "data", "dashboard.json")))
BG, PANEL, TXT, MUT, ACC = "#0a0f1c", "#111a2e", "#e6edf7", "#8a9bbb", "#4cc9f0"
COMM = ["#4cc9f0", "#f72585", "#ffd166", "#06d6a0", "#b388ff"]


def style(ax, fig):
    fig.patch.set_facecolor(BG); ax.set_facecolor(BG)
    for s in ax.spines.values(): s.set_color("#1e2a45")
    ax.tick_params(colors=MUT); ax.xaxis.label.set_color(MUT); ax.yaxis.label.set_color(MUT)


def charts():
    fig, ax = plt.subplots(figsize=(10, 6.2))
    for t in D["trajectories"]:
        p = t["p"]
        ax.plot([q[1] for q in p], [q[0] for q in p], lw=1.2 if t["hold"] else .5,
                color="#ff6b6b" if t["hold"] else ACC, alpha=.9 if t["hold"] else .25)
    for a in D["airports"]:
        ax.scatter(a["lon"], a["lat"], s=40 + a["pagerank"] * 3000, color=COMM[a["community"]], edgecolor="white", zorder=3)
        ax.text(a["lon"] + .3, a["lat"] + .2, a["code"], color=TXT, fontsize=8)
    style(ax, fig); ax.set_aspect("equal"); ax.set_xticks([]); ax.set_yticks([])
    fig.savefig(f"{OUT}/map.png", dpi=150, bbox_inches="tight"); plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 4.5))
    bh, bottom = D["band_hour"], [0] * 24
    for b, c, l in [("HIGH", ACC, "High (>8 km)"), ("MID", "#ffd166", "Medium"), ("LOW", "#06d6a0", "Low (<3 km)")]:
        ax.bar(range(24), bh[b], bottom=bottom, color=c, label=l); bottom = [x + y for x, y in zip(bottom, bh[b])]
    style(ax, fig); ax.set_xlabel("Hour of day (IST)"); ax.set_ylabel("plane positions"); ax.set_xticks(range(24))
    ax.legend(facecolor=BG, labelcolor=TXT, edgecolor="#1e2a45")
    fig.savefig(f"{OUT}/hourly.png", dpi=150, bbox_inches="tight"); plt.close(fig)

    fig, ax = plt.subplots(figsize=(8, 4.5))
    wr = D["worst_routes"][::-1]
    ax.barh([f"{r['o']} to {r['d']}" for r in wr], [r["ineff"] for r in wr], color="#ff9f1c")
    style(ax, fig); ax.set_xlabel("% extra distance flown")
    fig.savefig(f"{OUT}/routes.png", dpi=150, bbox_inches="tight"); plt.close(fig)


rgb = lambda h: RGBColor.from_string(h[1:])


def text(s, t, x, y, w, h, size, color, bold=False):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(size); p.font.color.rgb = rgb(color); p.font.bold = bold; p.space_after = Pt(10)


def stat(s, v, l, x, y, w=2.9):
    box = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(1.3))
    box.fill.solid(); box.fill.fore_color.rgb = rgb(PANEL); box.line.color.rgb = rgb("#1e2a45")
    text(s, v, x + .2, y + .12, w - .3, .6, 24, ACC, bold=True); text(s, l, x + .2, y + .75, w - .3, .5, 12, MUT)


def main():
    charts()
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    k = D["kpis"]
    pics = {"SkyTrace": ("map.png", 7.6, 1.0, 6.0), "What we found": ("hourly.png", 6.9, 1.8, 3.0),
            "The dashboard": ("dashboard.png", 7.0, 2.0, 3.45), "Airport network: groups and hubs": ("map.png", 8.6, 1.5, 5.6),
            "Why do we need this?": ("routes.png", 7.6, 4.3, 2.9)}
    for sd in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
        text(s, sd["kicker"], .6, .35, 8, .4, 12, ACC, bold=True)
        text(s, sd["title"], .6, .7, 12, .9, 34 if sd["title"] != "SkyTrace" else 54, TXT, bold=True)
        pic = pics.get(sd["title"])
        w = 6.6 if pic else 12
        y = 1.8
        if sd.get("sub"):
            text(s, sd["sub"], .6, 1.75, 6.8, 1, 22, TXT); y = 3.0
        if sd["title"] == "What we found":
            for i, (v, l) in enumerate([(f"{k['avg_ineff']}%", "extra distance"), (f"{k['total_excess_km']:,} km", "extra flying / day"),
                                        (f"{k['total_co2_t']} t", "CO2 avoidable / day"), (str(k["holdings"]), "planes circled")]):
                stat(s, v, l, .6 + i * 3.1, 5.6)
            text(s, ["•  " + b for b in sd["bullets"]], .6, 1.8, 6.1, 3.6, 17, TXT)
        else:
            text(s, ["•  " + b for b in sd["bullets"]], .6, y, w, 5, 19 if not pic else 18, TXT if not sd.get("sub") else MUT)
        if pic:
            f, x, py, h = pic
            s.shapes.add_picture(f"{OUT}/{f}", Inches(x), Inches(py), height=Inches(h))
        s.notes_slide.notes_text_frame.text = sd["notes"]
    prs.save(f"{OUT}/SkyTrace_Presentation.pptx"); print("saved", len(SLIDES), "slides")


if __name__ == "__main__":
    main()
